"""
Template Loader module.
Parses all 21 Jumia Seller Center product creation templates in templates/jumia_templates/
and extracts primary categories, subcategories, and Jumia category IDs.
"""
import os
import glob
import json
import logging
from pathlib import Path
from typing import Dict, Any, List, Set
import openpyxl

logger = logging.getLogger(__name__)

TEMPLATES_DIR = Path(__file__).resolve().parent / "templates" / "jumia_templates"
TAXONOMY_CACHE_FILE = Path(__file__).resolve().parent / "categories_taxonomy.json"

class TemplateCategoryLoader:
    def __init__(self, templates_dir: Path = TEMPLATES_DIR):
        self.templates_dir = templates_dir
        self.taxonomy: Dict[str, Any] = {}

    def parse_all_templates(self) -> Dict[str, Any]:
        """
        Scan all Excel template files and extract category trees and IDs.
        """
        if TAXONOMY_CACHE_FILE.exists():
            try:
                with open(TAXONOMY_CACHE_FILE, "r", encoding="utf-8") as f:
                    self.taxonomy = json.load(f)
                    if self.taxonomy.get("categories"):
                        logger.info(f"Loaded {len(self.taxonomy['categories'])} categories from taxonomy cache.")
                        return self.taxonomy
            except Exception as e:
                logger.warning(f"Failed to read taxonomy cache: {e}")

        files = glob.glob(str(self.templates_dir / "*.xlsx"))
        logger.info(f"Parsing {len(files)} Jumia template files...")

        primary_categories: Dict[str, List[Dict[str, Any]]] = {}
        all_leaf_categories: List[Dict[str, Any]] = []

        for f_path in files:
            try:
                wb = openpyxl.load_workbook(f_path, data_only=True)
                # Locate Categories sheet
                cat_sheet_name = None
                for name in wb.sheetnames:
                    if "cat" in name.lower():
                        cat_sheet_name = name
                        break

                if not cat_sheet_name:
                    continue

                ws = wb[cat_sheet_name]
                for row in ws.iter_rows(values_only=True):
                    if not row or not row[0]:
                        continue
                    val = str(row[0]).strip()
                    if val.upper() == "CATEGORIES" or not val:
                        continue

                    # Jumia format: "1029448 - Automobile / Cars"
                    parts = val.split(" - ", 1)
                    if len(parts) == 2:
                        cat_id = parts[0].strip()
                        cat_path = parts[1].strip()
                    else:
                        cat_id = ""
                        cat_path = val

                    # Extract primary category e.g. "Automobile"
                    path_parts = [p.strip() for p in cat_path.split("/") if p.strip()]
                    if not path_parts:
                        continue

                    primary = path_parts[0]
                    if primary not in primary_categories:
                        primary_categories[primary] = []

                    record = {
                        "id": cat_id,
                        "full_path": cat_path,
                        "primary": primary,
                        "subcategories": path_parts[1:],
                        "leaf": path_parts[-1]
                    }

                    # Avoid duplicate paths
                    if not any(x["full_path"] == cat_path for x in primary_categories[primary]):
                        primary_categories[primary].append(record)
                        all_leaf_categories.append(record)

            except Exception as e:
                logger.warning(f"Error parsing template {os.path.basename(f_path)}: {e}")

        # Map primary categories to standard Jumia catalog URL slugs
        slug_map = {
            "Automobile": "automobile",
            "Baby Products": "baby-products",
            "Books, Movies and Music": "books-movies-music",
            "Computing": "computing",
            "Electronics": "electronics",
            "Garden & Outdoors": "patio-lawn-garden",
            "Health & Beauty": "health-beauty",
            "Home & Office": "home-office",
            "Musical Instruments": "musical-instruments",
            "Phones & Tablets": "phones-tablets",
            "Sporting Goods": "sporting-goods",
            "Toys & Games": "toys-games",
            "Services": "services"
        }

        self.taxonomy = {
            "categories": sorted(list(primary_categories.keys())),
            "primary_categories": primary_categories,
            "total_subcategories": len(all_leaf_categories),
            "slug_map": slug_map
        }

        # Cache to json
        try:
            with open(TAXONOMY_CACHE_FILE, "w", encoding="utf-8") as f:
                json.dump(self.taxonomy, f, indent=2)
            logger.info(f"Saved {len(self.taxonomy['categories'])} primary categories to {TAXONOMY_CACHE_FILE}")
        except Exception as e:
            logger.error(f"Failed to cache taxonomy: {e}")

        return self.taxonomy

    def get_primary_categories(self) -> List[str]:
        if not self.taxonomy:
            self.parse_all_templates()
        return self.taxonomy.get("categories", [])

    def get_slug_for_category(self, category_name: str) -> str:
        if not self.taxonomy:
            self.parse_all_templates()
        slug_map = self.taxonomy.get("slug_map", {})
        return slug_map.get(category_name, category_name.lower().replace(" & ", "-").replace(" ", "-"))

# Global singleton
global_template_loader = TemplateCategoryLoader()

if __name__ == "__main__":
    loader = TemplateCategoryLoader()
    tax = loader.parse_all_templates()
    print("Parsed Categories from 21 templates:")
    for cat in tax["categories"]:
        count = len(tax["primary_categories"].get(cat, []))
        print(f"• {cat} ({count} subcategories) -> Slug: {loader.get_slug_for_category(cat)}")
